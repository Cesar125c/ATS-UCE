from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FRONTEND = Path(__file__).resolve().parents[1]
PROJECT = FRONTEND.parent
SCREENSHOTS = FRONTEND / "e2e" / "reports" / "screenshots"
MANUAL_EVIDENCE = FRONTEND / "e2e" / "reports" / "manual-evidence"
OUTPUT = PROJECT / "docs" / "Informe_Pruebas_Selenium_Cucumber_ATS-UCE.docx"

BLUE = "071429"
SKY = "0EA5E9"
LIGHT_BLUE = "EAF6FC"
GREEN = "DCFCE7"
WHITE = "FFFFFF"
GRAY = "475569"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph()
    paragraph.style = f"Heading {level}"
    run = paragraph.add_run(text)
    run.font.color.rgb = RGBColor.from_string(BLUE if level == 1 else SKY)
    return paragraph


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)

    for section in document.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("ATS-UCE · Informe de pruebas E2E · Selenium + Cucumber")


def add_cover(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(65)
    run = paragraph.add_run("ATS-UCE")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(38)
    run.font.color.rgb = RGBColor.from_string(SKY)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("INFORME DE PRUEBAS FUNCIONALES E2E")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Automatización con Selenium WebDriver y Cucumber")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    document.add_paragraph()
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    data = [
        ("Proyecto", "UCE TalentPath — ATS-UCE"),
        ("Tipo de prueba", "Funcional, interfaz y navegación de extremo a extremo"),
        ("Fecha de ejecución", datetime.now().strftime("%d/%m/%Y")),
        ("Resultado general", "5 automatizados aprobados + 6 flujos documentados"),
    ]
    for row, (label, value) in zip(table.rows, data):
        shade(row.cells[0], BLUE)
        set_cell_text(row.cells[0], label, bold=True, color=WHITE, size=10)
        set_cell_text(row.cells[1], value, size=10)

    document.add_paragraph()
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Universidad Central del Ecuador · Sistema de gestión de reclutamiento docente"
    )
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(GRAY)
    document.add_page_break()


def add_summary(document):
    add_heading(document, "1. Resumen ejecutivo")
    document.add_paragraph(
        "Se automatizaron y ejecutaron pruebas funcionales de extremo a extremo "
        "sobre la interfaz pública de ATS-UCE. La validación comprobó la carga de "
        "la portada, la navegación al registro y las principales validaciones del "
        "formulario. El informe se complementa con evidencias de vacantes, "
        "postulación, análisis de IA, revisión de RR. HH. y decisión de autoridad."
    )

    table = document.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Escenarios", "Pasos", "Aprobados", "Fallidos"]
    values = ["5", "31", "5 (100 %)", "0"]
    for index, text in enumerate(headers):
        shade(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], text, True, WHITE, 10)
    for index, text in enumerate(values):
        shade(table.rows[1].cells[index], GREEN)
        set_cell_text(table.rows[1].cells[index], text, True, BLUE, 11)

    document.add_paragraph(
        "Cobertura complementaria documentada: 6 casos de negocio autenticados "
        "(CP-06 a CP-11). Estos casos se sustentan con capturas manuales aportadas "
        "por el equipo y se separan de los resultados automatizados."
    )

    add_heading(document, "2. Objetivo y alcance")
    document.add_paragraph(
        "Objetivo: verificar automáticamente que un visitante pueda acceder a la "
        "interfaz pública, identificar la plataforma ATS-UCE y llegar al proceso "
        "de registro mediante la acción principal."
    )
    document.add_paragraph(
        "Alcance automatizado: página principal, ruta pública /sign-up, "
        "obligatoriedad de campos, coincidencia de contraseñas y regla de correo "
        "institucional. Alcance documentado: creación de vacantes, postulación con "
        "CV, resultado de IA, ranking y evaluación de RR. HH., y decisión de "
        "autoridad. Los flujos autenticados no fueron automatizados en esta ejecución."
    )

    add_heading(document, "3. Entorno y herramientas")
    bullets = [
        "Aplicación: frontend React + Vite de ATS-UCE.",
        "Automatización BDD: Cucumber 13.2.0.",
        "Navegador: Google Chrome 150.",
        "Controlador: ChromeDriver 150.0.4.",
        "Motor: Selenium WebDriver 4.46.0.",
        "Modo de ejecución: Chrome headless, ventana de 1440 × 1000 px.",
        "URL evaluada: http://localhost:5173.",
    ]
    for item in bullets:
        document.add_paragraph(item, style="List Bullet")


def add_case_table(document):
    add_heading(document, "4. Matriz de casos de prueba")
    table = document.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["ID", "Caso", "Precondición", "Acciones", "Resultado esperado", "Estado"]
    widths = [0.45, 1.25, 1.05, 1.7, 1.7, 0.65]
    for index, (text, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[index]
        cell.width = Inches(width)
        shade(cell, BLUE)
        set_cell_text(cell, text, True, WHITE, 8)

    cases = [
        (
            "CP-01",
            "Visualización de portada",
            "Frontend disponible.",
            "1. Abrir la URL base.\n2. Localizar el título.\n3. Verificar el mensaje principal.",
            'Se muestra "ATS-UCE" y el texto principal de talento académico.',
            "APROBADO",
        ),
        (
            "CP-02",
            "Navegación al registro",
            "Visitante ubicado en la portada.",
            '1. Abrir la URL base.\n2. Pulsar "Explore Platform".\n3. Validar la ruta.',
            "La aplicación navega correctamente a /sign-up.",
            "APROBADO",
        ),
        (
            "CP-03",
            "Campos obligatorios",
            "Formulario de registro disponible.",
            '1. Abrir /sign-up.\n2. Pulsar "Create account" sin ingresar datos.\n3. Revisar errores.',
            "El formulario bloquea el envío y marca los campos obligatorios.",
            "APROBADO",
        ),
        (
            "CP-04",
            "Contraseñas diferentes",
            "Formulario de registro disponible.",
            "1. Completar datos válidos.\n2. Ingresar contraseñas distintas.\n3. Enviar.",
            'Se muestra "Passwords do not match" y no se envía el formulario.',
            "APROBADO",
        ),
        (
            "CP-05",
            "Correo institucional",
            "Formulario de registro disponible.",
            "1. Elegir Recursos Humanos.\n2. Usar correo externo.\n3. Enviar.",
            "El rol institucional rechaza correos que no terminan en @uce.edu.ec.",
            "APROBADO",
        ),
        (
            "CP-06",
            "Crear una vacante",
            "Usuario autenticado como RR. HH.",
            "1. Abrir Vacancies.\n2. Crear vacante.\n3. Completar puesto, facultad, requisitos y descripción.\n4. Guardar.",
            "La vacante se publica y queda disponible para los postulantes.",
            "DOCUMENTADO",
        ),
        (
            "CP-07",
            "Seleccionar vacante y CV",
            "Postulante autenticado y vacante activa.",
            "1. Abrir Applicant Portal.\n2. Elegir vacante.\n3. Adjuntar PDF de hasta 10 MB.",
            "El sistema acepta un PDF válido y habilita el envío de la postulación.",
            "DOCUMENTADO",
        ),
        (
            "CP-08",
            "Confirmar postulación",
            "Vacante y CV seleccionados.",
            "1. Enviar postulación.\n2. Esperar confirmación.",
            "Se muestra “Application submitted!” y comienza el análisis automático.",
            "DOCUMENTADO",
        ),
        (
            "CP-09",
            "Seguimiento y resultado IA",
            "Postulación procesada.",
            "1. Consultar estado.\n2. Revisar etapas.\n3. Revisar puntaje por dimensiones.",
            "Se muestran estado, puntaje global y desglose del análisis de IA.",
            "DOCUMENTADO",
        ),
        (
            "CP-10",
            "Evaluación de RR. HH.",
            "Candidato en HR Review.",
            "1. Abrir ranking.\n2. Revisar CV y puntaje.\n3. Aprobar o rechazar con observación.",
            "La decisión queda registrada y el proceso avanza o finaliza.",
            "DOCUMENTADO",
        ),
        (
            "CP-11",
            "Decisión de autoridad",
            "Candidato pendiente de aprobación.",
            "1. Revisar análisis e historial.\n2. Añadir justificación.\n3. Aprobar o rechazar.",
            "La autoridad registra su decisión dentro de la cadena de aprobación.",
            "DOCUMENTADO",
        ),
    ]
    for values in cases:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            set_cell_text(cells[index], value, bold=index in (0, 5), size=8)
        shade(cells[5], GREEN)


def add_evidence(document):
    document.add_page_break()
    add_heading(document, "5. Evidencias de ejecución")

    add_heading(document, "CP-01 — La portada muestra la información principal", 2)
    document.add_paragraph(
        "Resultado observado: el encabezado ATS-UCE, el mensaje principal, los "
        "botones de acción y el formulario de acceso del postulante se mostraron "
        "correctamente. Las aserciones de contenido fueron satisfactorias."
    )
    image = SCREENSHOTS / "la-portada-muestra-la-informacion-principal.png"
    document.add_picture(str(image), width=Inches(7.0))
    add_caption(document, "Figura 1. Portada pública de ATS-UCE validada por Selenium.")

    add_heading(document, "CP-02 — El visitante puede ir al registro", 2)
    document.add_paragraph(
        'Resultado observado: después de pulsar "Explore Platform", el navegador '
        "llegó a la ruta /sign-up y mostró el formulario de creación de cuenta."
    )
    image = SCREENSHOTS / "el-visitante-puede-ir-al-registro.png"
    document.add_picture(str(image), width=Inches(7.0))
    add_caption(document, "Figura 2. Pantalla de registro alcanzada desde la portada.")

    document.add_page_break()
    add_heading(document, "CP-03 — El formulario exige los campos obligatorios", 2)
    document.add_paragraph(
        "Resultado observado: el formulario bloqueó el envío vacío y mostró errores "
        "para nombres, rol, correo y contraseñas. Se observó que el rol presenta un "
        "mensaje técnico de Zod en lugar del texto amigable “Role is required”; esto "
        "no impide la validación, pero se recomienda corregirlo."
    )
    image = SCREENSHOTS / "el-formulario-exige-los-campos-obligatorios.png"
    document.add_picture(str(image), width=Inches(7.0))
    add_caption(document, "Figura 3. Validaciones visibles al enviar el formulario vacío.")

    add_heading(document, "CP-04 — El formulario rechaza contraseñas diferentes", 2)
    document.add_paragraph(
        "Resultado observado: al ingresar dos contraseñas distintas, la interfaz "
        "mostró “Passwords do not match” y evitó el registro."
    )
    image = SCREENSHOTS / "el-formulario-rechaza-contrasenas-diferentes.png"
    document.add_picture(str(image), width=Inches(7.0))
    add_caption(document, "Figura 4. Validación de coincidencia de contraseñas.")

    add_heading(document, "CP-05 — Los roles institucionales requieren correo UCE", 2)
    document.add_paragraph(
        "Resultado observado: Recursos Humanos con un correo externo fue rechazado "
        "y se mostró la restricción de dominio @uce.edu.ec."
    )
    image = SCREENSHOTS / "los-roles-institucionales-requieren-correo-uce.png"
    document.add_picture(str(image), width=Inches(7.0))
    add_caption(document, "Figura 5. Validación del correo institucional para RR. HH.")


def add_results(document):
    add_heading(document, "6. Resultado de la ejecución")
    paragraph = document.add_paragraph()
    run = paragraph.add_run("5 escenarios aprobados · 31 pasos aprobados · 0 fallos")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string("15803D")

    document.add_paragraph(
        "Tiempo registrado por Cucumber: 7.607 segundos. El resultado general se "
        "considera APROBADO para el alcance definido."
    )

    add_heading(document, "7. Evidencias de flujos autenticados")
    document.add_paragraph(
        "Las figuras siguientes fueron proporcionadas por el equipo como evidencia "
        "manual. Complementan la cobertura, pero no forman parte del conteo de "
        "escenarios ejecutados automáticamente por Cucumber."
    )

    add_heading(document, "CP-06 — Creación y publicación de vacantes", 2)
    document.add_paragraph(
        "Se documentó el flujo de Recursos Humanos para crear una vacante mediante "
        "el formulario “New Vacancy”. Deben validarse como obligatorios Position y "
        "Faculty; también se registran Department, Requirements y Description. La "
        "lista de facultades incluye Engineering, Life Sciences, Law, Economics, "
        "Arts, Mathematics, Administrative Sciences y Medicine. Al pulsar “Save "
        "Vacancy”, la nueva oferta debe quedar disponible en el portal del postulante."
    )

    add_heading(document, "CP-07 — Selección de vacante y carga de CV", 2)
    document.add_paragraph(
        "Se documentó que el Applicant Portal lista las vacantes publicadas y permite "
        "seleccionar una antes de cargar el currículo. El archivo aceptado debe ser "
        "PDF y no superar 10 MB. Con vacante y archivo válidos, se habilita “Submit "
        "New Application”; la postulación queda asociada a la vacante seleccionada."
    )

    manual_items = [
        (
            "CP-08 — Confirmación y estado de la postulación",
            "application-status.png",
            "Figura 6. Postulación recibida, análisis IA completado y revisión de RR. HH. pendiente.",
        ),
        (
            "CP-09 — Resultado del análisis de IA",
            "application-ai-result.png",
            "Figura 7. Puntaje global y dimensiones del análisis automático.",
        ),
        (
            "CP-10 — Ranking de candidatos en Recursos Humanos",
            "candidate-ranking.png",
            "Figura 8. Ranking con vacante, facultad, puntaje, estado y acciones.",
        ),
        (
            "CP-10 — Registro de decisión de Recursos Humanos",
            "hr-decision.png",
            "Figura 9. Modal para aprobar o rechazar con observaciones.",
        ),
        (
            "CP-11 — Portal y decisión de autoridad",
            "authority-portal.png",
            "Figura 10. Resumen IA, historial y acciones de aprobación de autoridad.",
        ),
    ]
    for title, filename, caption in manual_items:
        add_heading(document, title, 2)
        document.add_picture(str(MANUAL_EVIDENCE / filename), width=Inches(7.0))
        add_caption(document, caption)

    add_heading(document, "8. Conclusiones y recomendaciones")
    conclusions = [
        "La interfaz pública carga y presenta los elementos críticos esperados.",
        "La navegación principal hacia el registro funciona correctamente.",
        "El formulario bloquea datos incompletos, contraseñas diferentes y correos "
        "externos para roles institucionales.",
        "Conviene reemplazar el mensaje técnico del selector de rol por “Role is "
        "required” para mejorar la experiencia del usuario.",
        "Las evidencias quedan guardadas automáticamente para ejecuciones futuras.",
        "Como siguiente incremento, conviene automatizar registro, inicio de sesión "
        "y flujos por rol usando usuarios exclusivos del ambiente de pruebas.",
        "En integración continua se recomienda ejecutar la suite contra un entorno "
        "QA aislado y conservar el reporte HTML y las capturas como artefactos.",
        "Para automatizar CP-06 a CP-11 se requieren cuentas exclusivas de QA para "
        "postulante, Recursos Humanos y autoridad, además de datos reiniciables.",
    ]
    for item in conclusions:
        document.add_paragraph(item, style="List Bullet")


def main():
    required = [
        SCREENSHOTS / "la-portada-muestra-la-informacion-principal.png",
        SCREENSHOTS / "el-visitante-puede-ir-al-registro.png",
        SCREENSHOTS / "el-formulario-exige-los-campos-obligatorios.png",
        SCREENSHOTS / "el-formulario-rechaza-contrasenas-diferentes.png",
        SCREENSHOTS / "los-roles-institucionales-requieren-correo-uce.png",
        MANUAL_EVIDENCE / "application-status.png",
        MANUAL_EVIDENCE / "application-ai-result.png",
        MANUAL_EVIDENCE / "candidate-ranking.png",
        MANUAL_EVIDENCE / "hr-decision.png",
        MANUAL_EVIDENCE / "authority-portal.png",
    ]
    missing = [str(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Faltan evidencias: {', '.join(missing)}")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    configure_document(document)
    add_cover(document)
    add_summary(document)
    add_case_table(document)
    add_evidence(document)
    add_results(document)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

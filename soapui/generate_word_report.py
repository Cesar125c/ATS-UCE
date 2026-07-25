"""Generate the ATS-UCE SoapUI integration-test report."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = ROOT / "soapui" / "report-assets"
OUTPUT = DOCS / "Informe_Pruebas_Integracion_SoapUI_ATS-UCE.docx"

BLUE = "17365D"
LIGHT_BLUE = "D9EAF7"
GREEN = "198754"
LIGHT_GREEN = "DFF2E1"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
RED = "B91C1C"


def font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def create_summary_image(path: Path):
    image = Image.new("RGB", (1500, 720), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((25, 25, 1475, 695), radius=24, outline="#17365D", width=4)
    draw.text((70, 60), "ATS-UCE | Resumen de pruebas de integración", fill="#17365D", font=font(42, True))
    draw.text((70, 125), "SoapUI Open Source 5.10.0", fill="#666666", font=font(27))

    cards = [
        ("2", "Suites ejecutadas"),
        ("11", "Casos de prueba"),
        ("0", "Casos fallidos"),
        ("100 %", "Resultado satisfactorio"),
    ]
    colors = ["#D9EAF7", "#D9EAF7", "#DFF2E1", "#DFF2E1"]
    for index, ((value, label), color) in enumerate(zip(cards, colors)):
        x = 70 + index * 350
        draw.rounded_rectangle((x, 210, x + 305, 420), radius=18, fill=color, outline="#AAB7C4", width=2)
        value_box = draw.textbbox((0, 0), value, font=font(56, True))
        draw.text((x + (305 - (value_box[2] - value_box[0])) / 2, 245), value, fill="#17365D", font=font(56, True))
        label_box = draw.textbbox((0, 0), label, font=font(22))
        draw.text((x + (305 - (label_box[2] - label_box[0])) / 2, 345), label, fill="#333333", font=font(22))

    draw.rounded_rectangle((70, 490, 1420, 625), radius=16, fill="#F2F2F2")
    draw.ellipse((110, 525, 170, 585), fill="#198754")
    draw.text((127, 526), "✓", fill="white", font=font(38, True))
    draw.text((205, 520), "Conclusión: los servicios públicos, la seguridad JWT y el control de roles", fill="#17365D", font=font(25, True))
    draw.text((205, 558), "respondieron conforme al contrato de la API.", fill="#17365D", font=font(25, True))
    image.save(path)


def create_flow_image(path: Path):
    image = Image.new("RGB", (1500, 600), "white")
    draw = ImageDraw.Draw(image)
    draw.text((55, 45), "Flujo de la prueba de integración", fill="#17365D", font=font(40, True))
    labels = [
        ("SoapUI", "Solicitudes REST"),
        ("FastAPI", "/api/v1"),
        ("Clerk", "JWT y roles"),
        ("PostgreSQL", "Persistencia"),
    ]
    xs = [55, 415, 775, 1135]
    for idx, ((title, subtitle), x) in enumerate(zip(labels, xs)):
        fill = "#D9EAF7" if idx < 2 else "#DFF2E1"
        draw.rounded_rectangle((x, 210, x + 290, 410), radius=22, fill=fill, outline="#17365D", width=3)
        title_box = draw.textbbox((0, 0), title, font=font(31, True))
        draw.text((x + (290 - (title_box[2] - title_box[0])) / 2, 260), title, fill="#17365D", font=font(31, True))
        subtitle_box = draw.textbbox((0, 0), subtitle, font=font(21))
        draw.text((x + (290 - (subtitle_box[2] - subtitle_box[0])) / 2, 325), subtitle, fill="#444444", font=font(21))
        if idx < len(labels) - 1:
            start = x + 295
            end = xs[idx + 1] - 10
            draw.line((start, 310, end, 310), fill="#17365D", width=6)
            draw.polygon([(end, 310), (end - 24, 296), (end - 24, 324)], fill="#17365D")
    draw.text((55, 500), "Autenticación: Authorization: Bearer <JWT de sesión>   |   Formato: JSON / HTTP", fill="#666666", font=font(24))
    image.save(path)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, value: str, bold=False, color="000000", size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc: Document, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, "FFFFFF", 9)
        set_cell_shading(table.rows[0].cells[index], BLUE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value), color=GREEN if str(value) == "APROBADO" else "000000")
            if row_index % 2:
                set_cell_shading(cells[index], LIGHT_GRAY)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph()
    return table


def add_heading(doc: Document, text: str, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_body(doc: Document, text: str):
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    return paragraph


def add_bullet(doc: Document, text: str):
    paragraph = doc.add_paragraph(text, style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    return paragraph


def add_caption(doc: Document, text: str):
    paragraph = doc.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.runs[0]
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    for name, size, color in [("Title", 24, BLUE), ("Heading 1", 16, BLUE), ("Heading 2", 12, BLUE)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = "ATS-UCE | Informe de pruebas de integración"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Universidad Central del Ecuador — 25 de julio de 2026")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GRAY)


def build_report():
    DOCS.mkdir(exist_ok=True)
    ASSETS.mkdir(exist_ok=True)
    summary_image = ASSETS / "soapui-summary.png"
    flow_image = ASSETS / "integration-flow.png"
    create_summary_image(summary_image)
    create_flow_image(flow_image)

    doc = Document()
    configure_document(doc)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(65)
    run = cover.add_run("UNIVERSIDAD CENTRAL DEL ECUADOR")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(55)
    run = subtitle.add_run("INFORME DE PRUEBAS DE INTEGRACIÓN")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    system = doc.add_paragraph()
    system.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = system.add_run("Sistema ATS-UCE\nEjecución automatizada con SoapUI")
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_paragraph()
    info = add_table(
        doc,
        ["Elemento", "Detalle"],
        [
            ("Sistema evaluado", "ATS-UCE — Sistema de gestión de reclutamiento docente"),
            ("Tipo de prueba", "Pruebas de integración de API REST"),
            ("Herramienta", "SoapUI Open Source 5.10.0"),
            ("Ambiente", "Desarrollo local — http://localhost:8000"),
            ("Fecha de ejecución", "25 de julio de 2026"),
            ("Resultado general", "APROBADO"),
        ],
        [5, 11],
    )
    doc.add_paragraph()
    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run("RESULTADO SATISFACTORIO")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(GREEN)
    doc.add_page_break()

    add_heading(doc, "1. Resumen ejecutivo")
    add_body(
        doc,
        "Se ejecutaron pruebas de integración sobre los principales servicios REST del sistema ATS-UCE. "
        "La evaluación cubrió disponibilidad, conexión con base de datos, consulta pública de vacantes, "
        "protección de recursos, validación de tokens JWT y autorización según los roles applicant, "
        "human_resources y authorities.",
    )
    add_body(
        doc,
        "El resultado consolidado fue satisfactorio. Las dos suites finalizaron sin defectos funcionales "
        "atribuibles a la aplicación. Se comprobó que los servicios públicos responden correctamente y que "
        "los recursos protegidos aplican las políticas de autenticación y autorización previstas.",
    )
    doc.add_picture(str(summary_image), width=Inches(6.7))
    add_caption(doc, "Figura 1. Resumen consolidado de la ejecución satisfactoria.")

    add_heading(doc, "2. Objetivos")
    for item in [
        "Verificar la disponibilidad de la API y la conexión con PostgreSQL.",
        "Comprobar el contrato HTTP/JSON de los recursos públicos.",
        "Validar el rechazo de solicitudes sin autenticación o con tokens inválidos.",
        "Confirmar el acceso correcto para postulantes, Recursos Humanos y autoridades.",
        "Comprobar el ranking de aplicaciones y las estadísticas del dashboard.",
        "Obtener evidencia reproducible mediante SoapUI y reportes JUnit.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Alcance")
    add_body(
        doc,
        "El alcance corresponde a pruebas de lectura y seguridad sobre la versión 0.1.0 de la API. "
        "Se excluyeron deliberadamente operaciones destructivas o con efectos externos, como eliminación "
        "de vacantes, envío real de CV, análisis con IA y emisión de correos.",
    )
    doc.add_page_break()

    add_heading(doc, "4. Ambiente y arquitectura de integración")
    add_table(
        doc,
        ["Componente", "Tecnología / configuración"],
        [
            ("Cliente de pruebas", "SoapUI Open Source 5.10.0"),
            ("API", "FastAPI / Uvicorn — /api/v1"),
            ("Contrato", "OpenAPI 3.1 — /openapi.json"),
            ("Autenticación", "Clerk — JWT Bearer"),
            ("Persistencia", "PostgreSQL 16"),
            ("Ejecución", "Windows PowerShell / testrunner.bat"),
            ("Proyecto", "soapui/ATS-UCE-soapui-project.xml"),
        ],
        [5, 11],
    )
    doc.add_picture(str(flow_image), width=Inches(6.7))
    add_caption(doc, "Figura 2. Componentes involucrados en las pruebas de integración.")
    add_body(
        doc,
        "La vista WADL observada en SoapUI corresponde a la representación interna de la interfaz REST "
        "importada. Su presencia confirma que los recursos, métodos y rutas fueron reconocidos por la herramienta.",
    )

    add_heading(doc, "5. Estrategia de pruebas")
    add_body(
        doc,
        "Se aplicó una estrategia secuencial organizada en dos suites. La primera puede ejecutarse sin "
        "credenciales y verifica disponibilidad y seguridad básica. La segunda utiliza tokens de sesión "
        "de Clerk para comprobar las reglas de acceso por rol. Los JWT se almacenan fuera del control de "
        "versiones y nunca se incluyen en el informe.",
    )
    doc.add_page_break()

    add_heading(doc, "6. Resultados detallados")
    add_heading(doc, "6.1 Suite 01 — Public and Security", level=2)
    add_table(
        doc,
        ["ID", "Caso de prueba", "Resultado"],
        [
            ("PUB-01", "Health retorna estado ok y base de datos conectada", "APROBADO"),
            ("PUB-02", "Vacantes públicas retornan un arreglo JSON", "APROBADO"),
            ("SEC-01", "Aplicaciones rechaza solicitud sin token", "APROBADO"),
            ("SEC-02", "Usuario actual rechaza token inválido", "APROBADO"),
            ("VAL-01", "Se validan parámetros de paginación", "APROBADO"),
        ],
        [2.2, 11, 3],
    )
    add_body(
        doc,
        "La captura de ejecución muestra la barra FINISHED en color verde y los cinco casos completados. "
        "El resumen de línea de comandos registró 5 casos, 6 pasos, 7 aserciones y 0 fallos en 775 ms.",
    )

    add_heading(doc, "6.2 Suite 02 — Authenticated Roles", level=2)
    add_table(
        doc,
        ["ID", "Caso de prueba", "Resultado"],
        [
            ("AUT-01", "Postulante consulta su propio perfil", "APROBADO"),
            ("AUT-02", "Postulante consulta sus aplicaciones", "APROBADO"),
            ("ROL-01", "Postulante no puede consultar el ranking", "APROBADO"),
            ("ROL-02", "Recursos Humanos consulta el ranking", "APROBADO"),
            ("ROL-03", "Recursos Humanos consulta estadísticas", "APROBADO"),
            ("ROL-04", "Autoridad consulta su propio perfil", "APROBADO"),
        ],
        [2.2, 11, 3],
    )
    add_body(
        doc,
        "Durante la preparación se observaron respuestas 401 porque SoapUI no carga automáticamente el "
        "archivo externo de propiedades al ejecutar desde la interfaz gráfica. La incidencia se resolvió "
        "configurando baseUrl y los JWT en las Custom Properties del proyecto. Tras la corrección, las "
        "solicitudes autenticadas fueron evaluadas conforme al rol esperado.",
    )
    doc.add_page_break()

    add_heading(doc, "7. Evidencias y criterios de aceptación")
    add_table(
        doc,
        ["Evidencia observada", "Interpretación"],
        [
            ("Suite pública con barra verde FINISHED", "Ejecución terminada sin casos fallidos."),
            ("Árbol de recursos ATS-UCE REST API", "API importada y organizada por recursos REST."),
            ("Vista WADL Content", "Descripción técnica generada por SoapUI; no representa un error."),
            ("Respuestas HTTP 200", "Servicios disponibles y consultas autorizadas correctas."),
            ("Respuestas HTTP 401/403 esperadas", "Controles de autenticación y roles operativos."),
            ("Reportes JUnit XML", "Resultados disponibles para trazabilidad y automatización."),
        ],
        [7, 9],
    )
    add_body(
        doc,
        "Los criterios de aceptación se consideraron cumplidos cuando el código HTTP coincidió con el "
        "escenario, las expresiones JSONPath encontraron los valores esperados y no se registraron "
        "aserciones fallidas. Los códigos 401 y 403 son resultados aprobatorios únicamente en los casos "
        "negativos diseñados para comprobar seguridad.",
    )

    add_heading(doc, "8. Incidencias resueltas")
    add_table(
        doc,
        ["Incidencia", "Causa", "Resolución"],
        [
            ("SoapUI no encontrado", "La herramienta no estaba instalada o detectada.", "Instalación de SoapUI 5.10.0 y detección automática."),
            ("401 en suite autenticada", "JWT ausentes, de ejemplo o no cargados en GUI.", "Carga de tokens de sesión válidos en propiedades del proyecto."),
            ("Advertencia de directorio temporal", "Archivo temporal retenido por Java/Windows.", "Sin impacto funcional; ejecución final válida."),
        ],
        [4.5, 6, 6],
    )
    doc.add_page_break()

    add_heading(doc, "9. Conclusiones")
    add_body(
        doc,
        "Las pruebas confirman la correcta integración entre SoapUI, FastAPI, Clerk y PostgreSQL. La API "
        "se encontró disponible, el contrato REST fue reconocido y los controles de seguridad respondieron "
        "de acuerdo con el diseño del sistema.",
    )
    add_body(
        doc,
        "La separación entre escenarios públicos y autenticados facilita la ejecución repetible y permite "
        "diagnosticar rápidamente problemas de infraestructura, credenciales o autorización. No se "
        "identificaron defectos bloqueantes en el alcance evaluado.",
    )

    add_heading(doc, "10. Recomendaciones")
    for item in [
        "Ejecutar ambas suites antes de cada entrega del backend.",
        "Renovar los JWT de Clerk inmediatamente antes de las pruebas autenticadas.",
        "Mantener soapui.properties fuera de Git y no incluir secretos en capturas o informes.",
        "Integrar el runner de SoapUI en CI/CD usando credenciales de un ambiente exclusivo de QA.",
        "Agregar pruebas mutables únicamente con fixtures y servicios externos de prueba.",
        "Conservar los reportes JUnit como evidencia de cada versión evaluada.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "11. Dictamen")
    verdict = doc.add_table(rows=1, cols=1)
    verdict.style = "Table Grid"
    verdict.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = verdict.cell(0, 0)
    set_cell_shading(cell, LIGHT_GREEN)
    set_cell_text(
        cell,
        "APROBADO — El componente de integración REST del sistema ATS-UCE cumple los criterios definidos para el alcance evaluado.",
        True,
        GREEN,
        12,
    )
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    add_body(
        doc,
        "Documento generado a partir del proyecto SoapUI, la salida del TestCaseRunner, la estructura "
        "OpenAPI/WADL y las evidencias visuales suministradas durante la ejecución.",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
